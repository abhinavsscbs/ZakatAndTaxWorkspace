import io
import json
import sys
from pathlib import Path
import pytest
sys.path.insert(0, str(Path(__file__).resolve().parents[1]/'backend'))
from app import app
from providers import AIClient, ProviderError

@pytest.fixture
def client(tmp_path, monkeypatch):
    monkeypatch.setenv('DATA_DIR', str(tmp_path))
    from config import Config
    monkeypatch.setattr(Config, 'DICTIONARY_FILE', tmp_path/'dictionary.json')
    return app.test_client()

def upload(client, provider='openai', text=b'Article 1. The synthetic VAT rate is fifteen percent.'):
    return client.post('/api/docs', data={'doc_id':'policy','file':(io.BytesIO(text),'policy.txt')},headers={'X-AI-Provider':provider})

def test_missing_key_is_explicit(client,monkeypatch):
    monkeypatch.delenv('OPENAI_API_KEY',raising=False)
    response=upload(client)
    assert response.status_code==503
    assert response.json['telemetry']['steps'][-1]['status']=='error'

def test_provider_isolation_and_replace_delete(client,monkeypatch):
    calls=[]
    def embed(self,texts,query=False):
        from flask import g
        calls.append((g.provider,query,len(texts)))
        return [[1.,0.] for _ in texts]
    monkeypatch.setattr(AIClient,'embed',embed)
    monkeypatch.setattr(AIClient,'chat',lambda *a,**k:'The rate is fifteen percent [policy:0].')
    assert upload(client).status_code==200
    result=client.post('/api/chat',json={'query':'rate?'},headers={'X-AI-Provider':'cohere'})
    assert result.status_code==200
    assert ('cohere',False,1) in calls
    assert result.json['sources'][0]['doc_id']=='policy'
    assert result.json['telemetry']['provider']=='cohere'
    assert upload(client,text=b'Updated rate.').status_code==200
    assert client.delete('/api/docs/policy').status_code==200
    assert client.get('/api/docs').json['documents']==[]
    assert client.post('/api/chat',json={'query':'rate?'}).status_code==400

def test_input_limits_and_no_sources(client):
    assert client.get('/api/providers',headers={'X-AI-Provider':'invalid'}).status_code==400
    assert client.post('/api/chat',json={'query':'q'*8001}).status_code==400
    assert client.post('/api/chat',json={'query':'question'}).status_code==400
    assert client.post('/api/tax/analyze',data={'file':(io.BytesIO(b'x'*16001),'large.txt')}).status_code==400

def test_api_protocols_and_usage(client,monkeypatch):
    sent=[]
    monkeypatch.setenv('OPENAI_API_KEY','test-key')
    monkeypatch.setenv('COHERE_API_KEY','test-key')
    class Response:
        ok=True; status_code=200; headers={'x-request-id':'fake-id'}
        def __init__(self,data):self.data=data
        def json(self):return self.data
    def post(url,**kwargs):
        sent.append((url,kwargs['json']))
        if url.endswith('/embeddings'):
            return Response({'data':[{'index':0,'embedding':[1,0]}],'usage':{'prompt_tokens':4,'total_tokens':4}})
        if url.endswith('/embed'):
            return Response({'embeddings':{'float':[[1,0]]},'meta':{'billed_units':{'input_tokens':4}}})
        if url.endswith('/chat/completions'):
            return Response({'choices':[{'message':{'content':'Answer [policy:0]'},'finish_reason':'stop'}],'usage':{'prompt_tokens':10,'completion_tokens':5,'total_tokens':15}})
        return Response({'message':{'content':[{'text':'Answer [policy:0]'}]},'finish_reason':'COMPLETE','usage':{'tokens':{'input_tokens':10,'output_tokens':5}}})
    monkeypatch.setattr('providers.requests.post',post)
    for provider in ('openai','cohere'):
        assert upload(client,provider).status_code==200
        response=client.post('/api/chat',json={'query':'What is the rate?'},headers={'X-AI-Provider':provider})
        assert response.status_code==200
        assert any(step.get('usage') for step in response.json['telemetry']['steps'])
    assert any(payload.get('input_type')=='search_query' for _,payload in sent)
    assert all('Authorization' not in str(step) for step in response.json['telemetry']['steps'])

def test_retry_and_output_limit(client,monkeypatch):
    monkeypatch.setenv('OPENAI_API_KEY','test-key')
    monkeypatch.setattr('providers.time.sleep',lambda _:None)
    class Response:
        headers={'Retry-After':'0'}
        def __init__(self,status,data):self.status_code=status;self.ok=status==200;self.data=data
        def json(self):return self.data
    responses=[Response(429,{}),Response(200,{'data':[{'index':0,'embedding':[1,0]}],'usage':{'total_tokens':3}})]
    monkeypatch.setattr('providers.requests.post',lambda *a,**k:responses.pop(0))
    response=upload(client)
    assert response.status_code==200
    embedding=next(s for s in response.json['telemetry']['steps'] if s['name']=='document_embedding')
    assert embedding['attempts']==2
    def limited(*a,**k):
        if a[0].endswith('/embeddings'):
            return Response(200,{'data':[{'index':0,'embedding':[1,0]}]})
        return Response(200,{'choices':[{'message':{'content':'partial'},'finish_reason':'length'}]})
    monkeypatch.setattr('providers.requests.post',limited)
    response=client.post('/api/chat',json={'query':'rate?'})
    assert response.status_code==422
    assert response.json['telemetry']['steps'][-1]['finish_reason']=='length'


def test_assessment_pipeline(client,monkeypatch,tmp_path):
    import importlib.util
    import sys
    from providers import MODELS
    def embed(self,texts,query=False):return [[1.,0.] for _ in texts]
    def chat(self,system,prompt,**kwargs):
        if prompt.startswith('\nCompare the two contracts'):
            return json.dumps({'differences':[], 'overall_summary':'No differences found.'})
        if prompt.startswith('\nIdentify distinct clauses'):
            return '[]'
        return 'VAT analysis with synthetic rate 15 [benchmark-policy:0].'
    monkeypatch.setattr(AIClient,'embed',embed)
    monkeypatch.setattr(AIClient,'chat',chat)
    root=Path(__file__).resolve().parents[1]
    spec=importlib.util.spec_from_file_location('assess',root/'scripts/assess.py')
    module=importlib.util.module_from_spec(spec);spec.loader.exec_module(module)
    monkeypatch.setattr(sys,'argv',['assess','--repeats','1','--output',str(tmp_path/'results')])
    assert module.main()==0
    reports=list((tmp_path/'results').glob('*/raw.json'))
    report=json.loads(reports[0].read_text())
    assert len(report['results'])==12
    assert all(row['success'] for row in report['results'])
    assert {row['provider'] for row in report['results']}=={'openai','cohere'}
    assert (reports[0].parent/'blind_review.csv').exists()


def test_comparison_quote_validation_and_exports(client,monkeypatch,tmp_path):
    from config import Config
    monkeypatch.setattr(Config,'DEFAULT_OUTPUT_DIR',tmp_path/'exports')
    monkeypatch.setattr(AIClient,'chat',lambda *a,**k:json.dumps({'differences':[{'older_contract':'Old payment terms.','newer_contract':'New payment terms.','difference_summary':'Terms changed','impact_score':'Moderate'}], 'overall_summary':'Payment terms changed.'}))
    def compare():
        return client.post('/api/comparisons/run',data={'files':[(io.BytesIO(b'Old payment terms.'),'old.txt'),(io.BytesIO(b'New payment terms.'),'new.txt')],'pairs':json.dumps([{'older':'old.txt','newer':'new.txt'},{'older':'old.txt','newer':'new.txt'}])})
    response=compare()
    assert response.status_code==200
    result=response.json['results'][0]
    assert result['success']
    assert response.json['results'][1]['success']
    assert client.get(result['download_url']).status_code==200
    monkeypatch.setattr(AIClient,'chat',lambda *a,**k:json.dumps({'differences':[{'older_contract':'Invented clause','newer_contract':'New payment terms.','impact_score':'Moderate'}], 'overall_summary':'Changed.'}))
    assert not compare().json['results'][0]['success']

def test_reasoning_model_omits_temperature_and_surfaces_400(client,monkeypatch):
    from providers import MODELS
    from telemetry import begin
    from flask import g
    monkeypatch.setenv('OPENAI_API_KEY','test-key')
    monkeypatch.delenv('OPENAI_TEMPERATURE',raising=False)
    monkeypatch.setitem(MODELS,'openai',{'chat':'gpt-6.1-sol','embedding':'text-embedding-3-small'})
    sent=[]
    class Response:
        ok=True;status_code=200;headers={}
        def json(self):return {'choices':[{'message':{'content':'OK'},'finish_reason':'stop'}]}
    def post(*a,**k):sent.append(k['json']);return Response()
    monkeypatch.setattr('providers.requests.post',post)
    with app.test_request_context():
        begin();g.provider='openai'
        assert AIClient().chat('System','Question')=='OK'
        assert 'temperature' not in sent[0]
    class Bad:
        ok=False;status_code=400;headers={'x-request-id':'req_failed'}
        def json(self):return {'error':{'message':'Unsupported temperature value','code':'unsupported_value','param':'temperature'}}
    monkeypatch.setattr('providers.requests.post',lambda *a,**k:Bad())
    with app.test_request_context():
        begin();g.provider='openai'
        with pytest.raises(ProviderError,match='Unsupported temperature value') as caught:
            AIClient().chat('System','Question')
        assert caught.value.status==400
        assert g.steps[-1]['provider_error_param']=='temperature'
        assert g.steps[-1]['provider_request_id']=='req_failed'

@pytest.mark.parametrize('provider', ['openai','cohere'])
def test_grounded_tax_uses_only_selected_sources_and_exports(client,monkeypatch,provider):
    import pandas as pd
    prompts=[]
    clause='The customer must pay all applicable taxes on the consulting fees.'
    monkeypatch.setattr(AIClient,'embed',lambda self,texts,query=False:[[1.,0.] for _ in texts])
    for doc_id,text in [('selected-policy',b'Article 1: A synthetic selected rule applies.'),('excluded-policy',b'EXCLUDED SECRET RULE must not be used.')]:
        assert client.post('/api/docs',data={'doc_id':doc_id,'file':(io.BytesIO(text),'source.txt')},headers={'X-AI-Provider':provider}).status_code==200
    def chat(self,system,prompt,**kwargs):
        prompts.append(prompt)
        if kwargs.get('step_name')=='clause_extraction':
            return json.dumps([{'title':'Taxes','full_text':clause,'category':'VAT','summary':'Tax responsibility'}])
        if kwargs.get('step_name')=='grounded_clause_tax_analysis':
            assert 'Article 1: A synthetic selected rule applies.' in prompt
            assert 'EXCLUDED SECRET RULE' not in prompt
            return json.dumps({'analysis':'### Implications\nReview the selected rule [selected-policy:0].','risk_level':'Medium','evidence_status':'supported','source_refs':['selected-policy:0']})
        return 'Contract overview only.'
    monkeypatch.setattr(AIClient,'chat',chat)
    response=client.post('/api/tax/analyze',data={'file':(io.BytesIO(clause.encode()),'contract.txt'),'analysis_mode':'grounded','selected_doc_ids':json.dumps(['selected-policy'])},headers={'X-AI-Provider':provider})
    assert response.status_code==200,response.json
    result=response.json
    row=result['analysis_results'][0]
    assert row['sources_used']==1
    assert row['evidence_status']=='supported'
    assert {s['doc_id'] for s in row['tax_info']['source_chunks']}=={'selected-policy'}
    assert result['selected_doc_ids']==['selected-policy']
    assert any(step['name']=='clause_evidence_validation' for step in result['telemetry']['steps'])
    export=client.post('/api/tax/export/excel',json=result)
    assert export.status_code==200
    workbook=pd.ExcelFile(io.BytesIO(export.data))
    assert 'Source_Evidence' in workbook.sheet_names
    evidence=pd.read_excel(workbook,sheet_name='Source_Evidence')
    assert evidence.iloc[0]['Document_ID']=='selected-policy'


def test_grounded_tax_rejects_empty_or_deleted_selections_before_generation(client,monkeypatch):
    def unexpected(*a,**k):raise AssertionError('No generation should happen')
    monkeypatch.setattr(AIClient,'chat',unexpected)
    for selection in ([],['deleted-policy']):
        response=client.post('/api/tax/analyze',data={'file':(io.BytesIO(b'Contract text.'),'contract.txt'),'analysis_mode':'grounded','selected_doc_ids':json.dumps(selection)})
        assert response.status_code==400


def test_grounded_tax_rejects_invented_citations_and_marks_insufficient(client,monkeypatch):
    from services.tax_service import TaxService
    from telemetry import begin
    from flask import g
    monkeypatch.setattr('knowledge.retrieve',lambda *a,**k:[{'doc_id':'allowed','chunk':0,'text':'No relevant tax provisions.','score':.3}])
    payload={'analysis':'Unsupported rule [invented:0].','risk_level':'High','evidence_status':'supported','source_refs':['invented:0']}
    monkeypatch.setattr(AIClient,'chat',lambda *a,**k:json.dumps(payload))
    clause={'title':'Tax','category':'Tax','full_text':'Contract clause'}
    with app.test_request_context():
        begin();g.provider='openai'
        with pytest.raises(ValueError,match='outside the selected evidence'):
            TaxService().grounded_clause_analysis(clause,['allowed'])
        payload.update(analysis='The sources do not establish a treatment.',risk_level='High',evidence_status='insufficient',source_refs=[])
        result,sources,cited=TaxService().grounded_clause_analysis(clause,['allowed'])
        assert result['risk_level']=='Undetermined'
        assert cited==[]

@pytest.mark.parametrize('provider', ['openai','cohere'])
def test_exact_contract_redaction_before_all_outbound_calls(client,monkeypatch,provider):
    from services.dictionary_service import DictionaryService
    dictionary=DictionaryService()
    dictionary.upsert('Secret Company', '[PARTY_A]')
    dictionary.upsert('Secret', '[PRIVATE]')
    original='1. Secret Company pays SAR 100.\n\n2. SECRET COMPANY delivers services.\n  Exact spacing stays.\n'
    expected='1. [PARTY_A] pays SAR 100.\n\n2. [PARTY_A] delivers services.\n  Exact spacing stays.\n'
    captured=[]
    def embed(self,texts,query=False):
        captured.extend(texts)
        return [[1.,0.] for _ in texts]
    def chat(self,system,prompt,**kwargs):
        captured.append(system+'\n'+prompt)
        assert kwargs.get('step_name')=='grounded_clause_tax_analysis'
        return json.dumps({'analysis':'Check the reference [policy:0].','risk_level':'Medium','evidence_status':'supported','source_refs':['policy:0']})
    monkeypatch.setattr(AIClient,'embed',embed)
    monkeypatch.setattr(AIClient,'chat',chat)
    assert upload(client,provider,text=b'Applicable source policy.').status_code==200
    response=client.post('/api/tax/analyze',data={'file':(io.BytesIO(original.encode()),'Secret Company.txt'),'analysis_mode':'grounded','selected_doc_ids':'["policy"]'},headers={'X-AI-Provider':provider})
    assert response.status_code==200,response.json
    result=response.json
    assert result['contract_text']==expected
    assert ''.join(item['clause_data']['full_text'] for item in result['analysis_results'])==expected
    assert 'executive_summary' not in result and 'detailed_analysis' not in result
    assert result['replacement_info']['matches_replaced']==2
    assert result['file_name']=='[PARTY_A].txt'
    assert all('secret' not in text.lower() for text in captured)
    assert {step['name'] for step in result['telemetry']['steps']}.isdisjoint({'clause_extraction','executive_summary','detailed_analysis'})
    export=client.post('/api/tax/export/excel',json=result)
    import pandas as pd
    workbook=pd.ExcelFile(io.BytesIO(export.data))
    assert 'Contract_Text' in workbook.sheet_names
    assert pd.read_excel(workbook,sheet_name='Contract_Text').iloc[0,0]==expected


def test_dictionary_shared_sessions_single_pass_and_concurrent_updates(client,monkeypatch):
    from services.dictionary_service import DictionaryService
    from concurrent.futures import ThreadPoolExecutor
    first=app.test_client();second=app.test_client()
    assert first.post('/api/dictionary',json={'word':'Acme Holdings','replacement':'[ORG]'}).status_code==200
    assert second.get('/api/dictionary').json['Acme Holdings']=='[ORG]'
    assert DictionaryService().load()['Acme Holdings']=='[ORG]'
    with ThreadPoolExecutor(max_workers=4) as pool:
        list(pool.map(lambda i:DictionaryService().upsert(f'entity-{i}',f'[ENTITY_{i}]'),range(12)))
    data=DictionaryService().load()
    assert all(f'entity-{i}' in data for i in range(12))
    assert json.loads(DictionaryService().path.read_text())==data
    assert DictionaryService.replace('Acme Holdings and ACME and $Cash', {'Acme':'[SHORT]','Acme Holdings':'Acme','[SHORT]':'other','$Cash':r'\1'})==('Acme and [SHORT] and \\1',3)


def test_exact_sectioning_and_docx_preserve_order(tmp_path):
    from services.tax_service import TaxService
    from services.document_service import DocumentService
    from docx import Document
    text='Intro\n\n1. Payment\n  exact spacing\n'+('Long clause '*700)+'\nالمادة 2: شروط\n'
    sections=TaxService.contract_sections(text)
    assert ''.join(section['full_text'] for section in sections)==text
    assert all(len(section['full_text'])<=3000 for section in sections)
    document=Document();document.add_paragraph('Before  table')
    table=document.add_table(rows=1,cols=2)
    table.cell(0,0).text='Left';table.cell(0,1).text='Right'
    document.add_paragraph('After table')
    path=tmp_path/'ordered.docx';document.save(path)
    assert DocumentService.extract_text(str(path))=='Before  table\nLeft\tRight\nAfter table'

@pytest.mark.parametrize('extension', ['csv','xlsx'])
def test_bulk_dictionary_merge_replace_and_invalid_atomicity(client,extension):
    from openpyxl import Workbook
    client.post('/api/dictionary',json={'word':'Existing','replacement':'[OLD]'})
    def content(rows):
        if extension=='csv':
            import csv
            buffer=io.StringIO();csv.writer(buffer).writerows(rows)
            return buffer.getvalue().encode('utf-8-sig')
        book=Workbook()
        for row in rows:book.active.append(row)
        buffer=io.BytesIO();book.save(buffer);return buffer.getvalue()
    def send(rows,mode):
        return client.post('/api/dictionary/import',data={'mode':mode,'file':(io.BytesIO(content(rows)),f'dictionary.{extension}')})
    merged=send([['Word','Replacement'],['Secret Name','[NAME]'],['existing','[UPDATED]']], 'merge')
    assert merged.status_code==200
    assert merged.json['existing']=='[UPDATED]'
    assert 'Existing' not in merged.json
    before=client.get('/api/dictionary').json
    bad=send([['Word','Replacement'],['good','[GOOD]'],['missing','']], 'replace')
    assert bad.status_code==400
    assert {key:value for key,value in client.get('/api/dictionary').json.items() if key != 'telemetry'} == {key:value for key,value in before.items() if key != 'telemetry'}
    replaced=send([['Word','Replacement'],['Only','[ONLY]']], 'replace')
    assert replaced.status_code==200
    assert 'existing' not in replaced.json and replaced.json['Only']=='[ONLY]'
    assert send([['Word','Replacement']], 'replace').status_code==400
    assert send([['bad','headers'],['one','two']], 'replace').status_code==400
    assert send([['Word','Replacement'],['x','y']], 'invalid').status_code==400


def test_library_original_and_both_indexes_survive_fresh_process(client,monkeypatch,tmp_path):
    import subprocess,os
    from services.dictionary_service import DictionaryService
    monkeypatch.setattr(AIClient,'embed',lambda self,texts,query=False:[[1.,0.] for _ in texts])
    monkeypatch.setattr(AIClient,'chat',lambda *a,**k:'Persisted source [policy:0].')
    assert upload(client).status_code==200
    assert client.post('/api/chat',json={'query':'question'},headers={'X-AI-Provider':'cohere'}).status_code==200
    listing=app.test_client().get('/api/docs').json['documents'][0]
    assert {item['provider'] for item in listing['indexes']}=={'openai','cohere'}
    assert client.get('/api/docs/policy/download').data==b'Article 1. The synthetic VAT rate is fifteen percent.'
    code='''
import sys
sys.path.insert(0,'backend')
from app import app
from providers import AIClient

def embed(self,texts,query=False):
    assert query, 'Persisted document vectors should not be regenerated'
    return [[1.,0.]]
AIClient.embed=embed
AIClient.chat=lambda *a,**k:'Stored answer [policy:0].'
client=app.test_client()
assert len(client.get('/api/docs').json['documents'])==1
for provider in ('openai','cohere'):
    response=client.post('/api/chat',json={'query':'question'},headers={'X-AI-Provider':provider})
    assert response.status_code==200,response.json
assert client.get('/api/docs/policy/download').status_code==200
print('Fresh-process persistence passed')
'''
    result=subprocess.run([sys.executable,'-c',code],cwd=Path(__file__).resolve().parents[1],env={**os.environ,'DATA_DIR':str(tmp_path)},capture_output=True,text=True)
    assert result.returncode==0,result.stderr+result.stdout
